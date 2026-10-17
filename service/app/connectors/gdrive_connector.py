"""Google Drive connector: CSV files, Google Sheets (every tab) and Excel
workbooks (every sheet) in a Drive, read-only.

Two ways to sign in:
- "Sign in with Google" (OAuth 2.0, scope drive.readonly). The server keeps
  only the refresh token, encrypted; the browser never sees it. Needs a
  Google Cloud OAuth client: GOOGLE_OAUTH_CLIENT_ID / GOOGLE_OAUTH_CLIENT_SECRET.
- A service-account key (JSON). Share a Drive folder with the service
  account's e-mail address and paste the key. No browser redirect needed.

How rows are read: each file/tab is turned into CSV text and parsed by the
SAME CSVConnector the upload flow uses, so a table read from Drive is scored
exactly like the same table uploaded by hand (same type inference, same
null-sentinel handling, same checks).

Limits, stated plainly:
- At most MAX_FILES files are listed (most recently modified first); point
  the connection at a folder to narrow it down.
- CSV downloads stop after MAX_DOWNLOAD_BYTES; Excel files larger than that
  are refused. Only the first `row_limit` data rows of each table are read.
- The first row of each sheet is treated as the header row.
- Formulas are read as their displayed values.
"""
from __future__ import annotations

import base64
import csv
import datetime as _dt
import hashlib
import io
import json
import logging
import os
import re
import tempfile
import time
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlencode

from app.canonical.models import CanonicalTable
from app.connectors.base import ConnectorConfig
from app.connectors.source import (
    DEFAULT_ROW_LIMIT,
    ConnectorError,
    ExtractedTable,
    TableInfo,
    TabularSource,
    clamp_row_limit,
)

_log = logging.getLogger(__name__)

SCOPE = "https://www.googleapis.com/auth/drive.readonly"
AUTH_URL = "https://accounts.google.com/o/oauth2/v2/auth"
TOKEN_URL = "https://oauth2.googleapis.com/token"
REVOKE_URL = "https://oauth2.googleapis.com/revoke"
DRIVE = "https://www.googleapis.com/drive/v3"
SHEETS = "https://sheets.googleapis.com/v4"

MIME_CSV = "text/csv"
MIME_SHEET = "application/vnd.google-apps.spreadsheet"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_FOLDER = "application/vnd.google-apps.folder"

MAX_FILES = 100
MAX_DOWNLOAD_BYTES = 30 * 1024 * 1024
HTTP_TIMEOUT_S = 30

# Tests replace this with an httpx.MockTransport; production uses the network.
HTTP_TRANSPORT: Any = None


class DriveError(ConnectorError):
    pass


def oauth_client() -> tuple[str, str] | None:
    cid = os.environ.get("GOOGLE_OAUTH_CLIENT_ID", "").strip()
    secret = os.environ.get("GOOGLE_OAUTH_CLIENT_SECRET", "").strip()
    return (cid, secret) if cid and secret else None


def _http():
    import httpx

    return httpx.Client(timeout=HTTP_TIMEOUT_S, transport=HTTP_TRANSPORT, follow_redirects=True)


def parse_folder_id(value: str | None) -> str | None:
    """Accepts a folder ID or any Drive folder link."""
    v = (value or "").strip()
    if not v:
        return None
    m = re.search(r"/folders/([A-Za-z0-9_-]+)", v) or re.search(r"[?&]id=([A-Za-z0-9_-]+)", v)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{10,}", v):
        return v
    raise DriveError("That doesn't look like a Google Drive folder link or ID.")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, _dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    return str(value)


def _clean_header(raw: list[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw):
        name = _cell(h).strip() or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _rows_to_table(name: str, matrix: list[list[Any]], row_limit: int, cfg: ConnectorConfig) -> tuple[CanonicalTable, int, bool]:
    """Header + data rows -> CanonicalTable through the upload path's
    CSVConnector. Returns (table, rows_read, more_rows_exist)."""
    from app.connectors.csv_connector import CSVConnector

    while matrix and not any(_cell(c).strip() for c in matrix[-1]):
        matrix.pop()  # trailing blank rows in sheets are not data
    if not matrix:
        raise DriveError(f"'{name}' is empty.")
    header = _clean_header(matrix[0])
    width = len(header)
    data = [r for r in matrix[1:]]
    more = len(data) > row_limit
    data = data[:row_limit]
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    for r in data:
        cells = [_cell(c) for c in list(r)[:width]]
        w.writerow(cells + [""] * (width - len(cells)))
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "table.csv"
        path.write_text(buf.getvalue(), encoding="utf-8")
        table = CSVConnector(ConnectorConfig(
            tenant_id=cfg.tenant_id, source_id=cfg.source_id, settings={"file_path": str(path)},
        )).extract()[0]
    table.table_name = name
    return table, len(data), more


# ---- OAuth -------------------------------------------------------------------

def _oauth_provider():
    from app.connectors.registry import OAuthProvider

    class GoogleOAuth(OAuthProvider):
        when = ("auth_mode", "oauth")

        def configured(self) -> bool:
            return oauth_client() is not None

        def authorization_url(self, *, state: str, redirect_uri: str, code_challenge: str) -> str:
            client = oauth_client()
            if client is None:
                raise DriveError(
                    "Google sign-in is not set up on this server (GOOGLE_OAUTH_CLIENT_ID / "
                    "GOOGLE_OAUTH_CLIENT_SECRET). Use a service-account key instead, or ask the admin."
                )
            return AUTH_URL + "?" + urlencode({
                "client_id": client[0], "redirect_uri": redirect_uri, "response_type": "code",
                "scope": SCOPE, "access_type": "offline", "prompt": "consent",
                "include_granted_scopes": "true", "state": state,
                "code_challenge": code_challenge, "code_challenge_method": "S256",
            })

        def exchange(self, *, code: str, redirect_uri: str, code_verifier: str) -> dict[str, str]:
            client = oauth_client()
            if client is None:
                raise DriveError("Google sign-in is not set up on this server.")
            with _http() as http:
                r = http.post(TOKEN_URL, data={
                    "code": code, "client_id": client[0], "client_secret": client[1],
                    "redirect_uri": redirect_uri, "grant_type": "authorization_code",
                    "code_verifier": code_verifier,
                })
            if r.status_code != 200:
                raise DriveError("Google did not accept the sign-in. Please try again.")
            body = r.json()
            if not body.get("refresh_token"):
                raise DriveError("Google did not return a long-lived token. Remove this app's access at "
                                 "myaccount.google.com/permissions and sign in again.")
            granted = set(str(body.get("scope", SCOPE)).split())
            if SCOPE not in granted:
                raise DriveError("Drive read access was not granted. Tick the Drive box on Google's screen.")
            return {"refresh_token": body["refresh_token"]}

    return GoogleOAuth()


# ---- connector ---------------------------------------------------------------

class GoogleDriveConnector(TabularSource):
    """settings: auth_mode ("oauth" | "service_account"), folder (optional
    link or ID), refresh_token (oauth, secret), service_account_json
    (service_account, secret), row_limit (optional)."""

    def __init__(self, config: ConnectorConfig):
        super().__init__(config)
        self._token: str | None = None
        self._token_exp = 0.0
        self._account: str | None = None
        self._xlsx_cache: dict[str, bytes] = {}
        self._files: dict[str, dict] | None = None

    # -- auth --
    def _mode(self) -> str:
        return str(self.config.settings.get("auth_mode") or "oauth")

    def _access_token(self) -> str:
        if self._token and time.time() < self._token_exp - 60:
            return self._token
        s = self.config.settings
        if self._mode() == "service_account":
            token, exp = self._service_account_token(s.get("service_account_json"))
        else:
            token, exp = self._refresh(s.get("refresh_token"))
        self._token, self._token_exp = token, exp
        return token

    def _refresh(self, refresh_token: str | None) -> tuple[str, float]:
        client = oauth_client()
        if not refresh_token:
            raise DriveError("This connection has no Google sign-in yet. Use 'Sign in with Google'.")
        if client is None:
            raise DriveError("Google sign-in is not set up on this server.")
        with _http() as http:
            r = http.post(TOKEN_URL, data={
                "client_id": client[0], "client_secret": client[1],
                "refresh_token": refresh_token, "grant_type": "refresh_token",
            })
        if r.status_code != 200:
            raise DriveError("Google sign-in expired or was revoked. Disconnect and sign in again.")
        body = r.json()
        return body["access_token"], time.time() + float(body.get("expires_in", 3600))

    def _service_account_token(self, raw: str | None) -> tuple[str, float]:
        try:
            info = json.loads(raw or "")
            email, key = info["client_email"], info["private_key"]
            token_uri = info.get("token_uri") or TOKEN_URL
        except (ValueError, KeyError, TypeError):
            raise DriveError("The service-account key must be the JSON file Google gave you "
                             "(it has client_email and private_key).") from None
        if token_uri != TOKEN_URL:
            raise DriveError("The service-account key points at an unexpected token server.")
        try:
            from google.auth import crypt, jwt
        except ImportError:
            raise DriveError("Service-account sign-in needs the google-auth package on the server.") from None
        now = int(time.time())
        try:
            signer = crypt.RSASigner.from_service_account_info(info)
        except Exception:  # noqa: BLE001
            raise DriveError("The service-account private key could not be read.") from None
        assertion = jwt.encode(signer, {
            "iss": email, "scope": SCOPE, "aud": TOKEN_URL, "iat": now, "exp": now + 3600,
        })
        if isinstance(assertion, bytes):
            assertion = assertion.decode()
        with _http() as http:
            r = http.post(TOKEN_URL, data={
                "grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer", "assertion": assertion,
            })
        if r.status_code != 200:
            raise DriveError("Google rejected the service-account key (disabled, deleted, or wrong project).")
        body = r.json()
        self._account = email
        return body["access_token"], time.time() + float(body.get("expires_in", 3600))

    # -- HTTP --
    def _get(self, url: str, params: dict | None = None, *, stream_limit: int | None = None):
        import httpx

        headers = {"Authorization": f"Bearer {self._access_token()}"}
        try:
            with _http() as http:
                if stream_limit is None:
                    r = http.get(url, params=params, headers=headers)
                    self._raise_for(r)
                    return r.json()
                with http.stream("GET", url, params=params, headers=headers) as r:
                    if r.status_code >= 400:
                        r.read()
                        self._raise_for(r)
                    buf = bytearray()
                    truncated = False
                    for chunk in r.iter_bytes():
                        buf.extend(chunk)
                        if len(buf) > stream_limit:
                            truncated = True
                            break
                    return bytes(buf[:stream_limit]), truncated
        except httpx.TimeoutException:
            raise DriveError("Google Drive timed out. Try again, or pick fewer files.") from None
        except httpx.HTTPError:
            raise DriveError("Could not reach Google Drive.") from None

    @staticmethod
    def _raise_for(r) -> None:
        if r.status_code < 400:
            return
        if r.status_code == 401:
            raise DriveError("Google sign-in expired or was revoked. Disconnect and sign in again.")
        if r.status_code == 403:
            try:
                reason = r.json().get("error", {}).get("errors", [{}])[0].get("reason", "")
            except Exception:  # noqa: BLE001
                reason = ""
            if "RateLimit" in reason or "rateLimit" in reason or "quota" in reason.lower():
                raise DriveError("Google Drive rate limit reached. Wait a minute and try again.")
            if "accessNotConfigured" in reason or "SERVICE_DISABLED" in str(r.text):
                raise DriveError("The Google Drive or Sheets API is not enabled for this Google Cloud project.")
            raise DriveError("Google denied access to that file or folder (is it shared with this account?).")
        if r.status_code == 404:
            raise DriveError("That file or folder was not found (or isn't shared with this account).")
        if r.status_code == 429:
            raise DriveError("Google Drive rate limit reached. Wait a minute and try again.")
        raise DriveError(f"Google Drive returned an error (HTTP {r.status_code}).")

    # -- contract --
    def check(self) -> None:
        about = self._get(f"{DRIVE}/about", {"fields": "user(emailAddress)"})
        self._account = self._account or about.get("user", {}).get("emailAddress")
        folder = parse_folder_id(self.config.settings.get("folder"))
        if folder:
            meta = self._get(f"{DRIVE}/files/{quote(folder)}",
                             {"fields": "id,name,mimeType", "supportsAllDrives": "true"})
            if meta.get("mimeType") != MIME_FOLDER:
                raise DriveError("That link is a file, not a folder. Paste a folder link.")

    def account_label(self) -> str | None:
        return self._account

    def _list_files(self) -> dict[str, dict]:
        if self._files is not None:
            return self._files
        q = (f"trashed = false and (mimeType = '{MIME_CSV}' or mimeType = '{MIME_SHEET}' "
             f"or mimeType = '{MIME_XLSX}')")
        folder = parse_folder_id(self.config.settings.get("folder"))
        if folder:
            q += f" and '{folder}' in parents"
        files: list[dict] = []
        token = None
        while len(files) < MAX_FILES:
            params = {
                "q": q, "pageSize": str(min(100, MAX_FILES - len(files))),
                "fields": "nextPageToken,files(id,name,mimeType,size,modifiedTime)",
                "orderBy": "modifiedTime desc", "supportsAllDrives": "true",
                "includeItemsFromAllDrives": "true",
            }
            if token:
                params["pageToken"] = token
            page = self._get(f"{DRIVE}/files", params)
            files.extend(page.get("files", []))
            token = page.get("nextPageToken")
            if not token:
                break
        self._files = {f["id"]: f for f in files}
        return self._files

    def _xlsx_bytes(self, file_id: str, size: int | None) -> bytes:
        if file_id in self._xlsx_cache:
            return self._xlsx_cache[file_id]
        if size is not None and size > MAX_DOWNLOAD_BYTES:
            raise DriveError(f"That Excel file is larger than {MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB.")
        data, truncated = self._get(f"{DRIVE}/files/{quote(file_id)}",
                                    {"alt": "media", "supportsAllDrives": "true"},
                                    stream_limit=MAX_DOWNLOAD_BYTES)
        if truncated:
            raise DriveError(f"That Excel file is larger than {MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB.")
        self._xlsx_cache[file_id] = data
        return data

    @staticmethod
    def _open_xlsx(data: bytes):
        try:
            import openpyxl
        except ImportError:
            raise DriveError("Reading Excel files needs the openpyxl package on the server.") from None
        try:
            return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            raise DriveError("That Excel file could not be opened (corrupt or password-protected).") from None

    def list_tables(self) -> list[TableInfo]:
        out: list[TableInfo] = []
        for f in self._list_files().values():
            stem = re.sub(r"\.(csv|xlsx)$", "", f["name"], flags=re.I)
            mime = f["mimeType"]
            if mime == MIME_CSV:
                out.append(TableInfo(stem, None, [], [], id=f"csv:{f['id']}",
                                     description=f"CSV file · {f['name']}"))
            elif mime == MIME_SHEET:
                meta = self._get(f"{SHEETS}/spreadsheets/{quote(f['id'])}",
                                 {"fields": "sheets.properties(sheetId,title,sheetType)"})
                tabs = [s["properties"] for s in meta.get("sheets", [])
                        if s.get("properties", {}).get("sheetType", "GRID") == "GRID"]
                for t in tabs:
                    name = stem if len(tabs) == 1 else t["title"]
                    out.append(TableInfo(name, None, [], [], id=f"gsheet:{f['id']}:{t['sheetId']}",
                                         description=f"Google Sheet · {f['name']} › {t['title']}"))
            elif mime == MIME_XLSX:
                try:
                    wb = self._open_xlsx(self._xlsx_bytes(f["id"], int(f["size"]) if f.get("size") else None))
                except DriveError as exc:
                    _log.info("skipping xlsx %s: %s", f["id"], exc)
                    continue
                names = list(wb.sheetnames)
                wb.close()
                for i, title in enumerate(names):
                    name = stem if len(names) == 1 else title
                    out.append(TableInfo(name, None, [], [], id=f"xlsx:{f['id']}:{i}",
                                         description=f"Excel · {f['name']} › {title}"))
        # Names drive cross-table matching (account_id -> accounts), so keep
        # them as-is and only suffix true duplicates.
        seen: dict[str, int] = {}
        for t in out:
            if t.name in seen:
                seen[t.name] += 1
                t.name = f"{t.name}_{seen[t.name]}"
            else:
                seen[t.name] = 1
        return out

    def extract_selected(self, table_ids: list[str], row_limit: int | None = None) -> list[ExtractedTable]:
        limit = clamp_row_limit(row_limit or self.config.settings.get("row_limit"), DEFAULT_ROW_LIMIT)
        catalog = {t.id: t for t in self.list_tables()}
        unknown = [i for i in table_ids if i not in catalog]
        if unknown:
            raise DriveError(f"Not available in this Drive connection: {', '.join(unknown)}")
        results = []
        for tid in table_ids:
            info = catalog[tid]
            kind, file_id, *rest = tid.split(":")
            notes: list[str] = []
            if kind == "csv":
                data, truncated = self._get(f"{DRIVE}/files/{quote(file_id)}",
                                            {"alt": "media", "supportsAllDrives": "true"},
                                            stream_limit=MAX_DOWNLOAD_BYTES)
                text = data.decode("utf-8-sig", errors="replace")
                if truncated:
                    text = text[: text.rfind("\n")] if "\n" in text else ""
                    notes.append(f"'{info.name}': only the first {MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB "
                                 "of the file were read.")
                matrix = list(csv.reader(io.StringIO(text)))
            elif kind == "gsheet":
                sheet_id = rest[0]
                meta = self._get(f"{SHEETS}/spreadsheets/{quote(file_id)}",
                                 {"fields": "sheets.properties(sheetId,title)"})
                title = next((s["properties"]["title"] for s in meta.get("sheets", [])
                              if str(s["properties"]["sheetId"]) == sheet_id), None)
                if title is None:
                    raise DriveError(f"The tab behind '{info.name}' no longer exists.")
                rng = "'" + title.replace("'", "''") + f"'!1:{limit + 2}"
                vals = self._get(f"{SHEETS}/spreadsheets/{quote(file_id)}/values/{quote(rng, safe='')}",
                                 {"valueRenderOption": "FORMATTED_VALUE", "majorDimension": "ROWS"})
                matrix = vals.get("values", [])
            else:  # xlsx
                f = self._list_files()[file_id]
                wb = self._open_xlsx(self._xlsx_bytes(file_id, int(f["size"]) if f.get("size") else None))
                ws = wb.worksheets[int(rest[0])]
                matrix = [list(r) for r in ws.iter_rows(max_row=limit + 2, values_only=True)]
                wb.close()
            table, n, more = _rows_to_table(info.name, matrix, limit, self.config)
            results.append(ExtractedTable(
                table=table, rows_scanned=n, approx_total_rows=None, ordered_by_primary_key=False,
                skipped_columns={}, hit_row_cap=more, stable_row_order=True, notes=notes,
            ))
        return results

    def revoke(self) -> None:
        token = self.config.settings.get("refresh_token")
        if not token or self._mode() != "oauth":
            return
        try:
            with _http() as http:
                http.post(REVOKE_URL, data={"token": token})
        except Exception:  # noqa: BLE001 - best effort
            _log.info("google token revoke failed")


def pkce_pair() -> tuple[str, str]:
    verifier = base64.urlsafe_b64encode(os.urandom(48)).rstrip(b"=").decode()
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b"=").decode()
    return verifier, challenge


def _specs():
    from app.connectors.registry import ConnectorSpec, FieldSpec

    return [ConnectorSpec(
        key="google_drive", label="Google Drive", category="files",
        description="CSV files, Google Sheets (every tab) and Excel workbooks in your Drive. Read-only.",
        fields=(
            FieldSpec("auth_mode", "Sign in with", kind="select", default="oauth",
                      options=(("oauth", "My Google account"), ("service_account", "Service-account key (JSON)"))),
            FieldSpec("folder", "Folder link",
                      placeholder="https://drive.google.com/drive/folders/...",
                      help="Only files in this folder are listed. Leave blank for your most recent files."),
            FieldSpec("service_account_json", "Service-account key (JSON)", kind="textarea", secret=True,
                      required=True, show_if=(("auth_mode", "service_account"),),
                      help="Share the folder with the service account's e-mail first."),
        ),
        factory=GoogleDriveConnector,
        oauth=_oauth_provider(),
    )]


CONNECTOR_SPECS = _specs()
